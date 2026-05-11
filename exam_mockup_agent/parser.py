"""XLSX Parser - Parse and validate uploaded XLSX files into structured Question Bank."""

from openpyxl import load_workbook

from exam_mockup_agent.models import Question

# Original format columns
REQUIRED_COLUMNS = [
    "Question",
    "Option_A",
    "Option_B",
    "Option_C",
    "Option_D",
    "Correct_Answer",
    "Question_Type",
]

OPTIONAL_COLUMNS = ["Option_E", "Option_F"]

# Alternative format columns (e.g. Google Sheets exports)
ALT_REQUIRED_COLUMNS = [
    "Question Text",
    "Option A",
    "Option B",
    "Option C",
    "Option D",
    "Correct answer",
]

ALT_OPTIONAL_COLUMNS = ["Option E", "Option F", "Question Number"]


class ValidationError(Exception):
    """Raised when XLSX validation fails."""

    pass


def _detect_format(headers: list) -> str:
    """Detect which XLSX format is being used.

    Returns 'original' for the standard format or 'alt' for the alternative format.
    """
    if all(col in headers for col in REQUIRED_COLUMNS):
        return "original"
    if all(col in headers for col in ALT_REQUIRED_COLUMNS):
        return "alt"
    return "unknown"


def validate_xlsx_structure(file_path: str) -> tuple[bool, list[str]]:
    """Validate XLSX has required columns.

    Returns (is_valid, missing_columns).
    Supports both original and alternative column formats.
    """
    wb = load_workbook(file_path, read_only=True)
    ws = wb.active
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    wb.close()

    fmt = _detect_format(headers)
    if fmt == "original":
        missing = [col for col in REQUIRED_COLUMNS if col not in headers]
    elif fmt == "alt":
        missing = [col for col in ALT_REQUIRED_COLUMNS if col not in headers]
    else:
        # Try original format for error reporting
        missing = [col for col in REQUIRED_COLUMNS if col not in headers]

    return (len(missing) == 0, missing)


def _parse_alt_correct_answer(raw_value: str) -> tuple[list[str], str]:
    """Parse the alternative format's 'Correct answer' column.

    The value is typically 'B\\nExplanation...' or 'B, E\\nExplanation...'.
    Returns (correct_answers, question_type).
    """
    if not raw_value:
        return ([], "single")

    # Split on newline — answer letters are before the first newline
    parts = raw_value.split("\n", 1)
    answer_part = parts[0].strip()

    if "," in answer_part:
        correct_answers = [a.strip().upper() for a in answer_part.split(",")]
        question_type = "multiple"
    else:
        correct_answers = [answer_part.upper()]
        question_type = "single"

    return (correct_answers, question_type)


def parse_xlsx(file_path: str) -> list[Question]:
    """Parse XLSX file and return list of Question objects.

    Supports both original format (with Question_Type column) and
    alternative format (with 'Correct answer' containing letter + explanation).

    Raises ValidationError if file is invalid.
    """
    try:
        wb = load_workbook(file_path, read_only=True, data_only=True)
    except Exception as e:
        raise ValidationError(f"Cannot open XLSX file: {e}")

    ws = wb.active
    rows = list(ws.iter_rows(min_row=1, values_only=True))
    wb.close()

    if not rows:
        raise ValidationError("XLSX file is empty")

    headers = list(rows[0])
    fmt = _detect_format(headers)

    if fmt == "original":
        return _parse_original_format(headers, rows)
    elif fmt == "alt":
        return _parse_alt_format(headers, rows)
    else:
        # Fall back to original format validation for error message
        missing = [col for col in REQUIRED_COLUMNS if col not in headers]
        raise ValidationError(
            f"Missing required columns: {', '.join(missing)}"
        )


def _parse_original_format(headers: list, rows: list) -> list[Question]:
    """Parse the original XLSX format with explicit Question_Type column."""
    missing = [col for col in REQUIRED_COLUMNS if col not in headers]
    if missing:
        raise ValidationError(
            f"Missing required columns: {', '.join(missing)}"
        )

    col_index = {name: idx for idx, name in enumerate(headers)}
    questions = []

    for row_num, row in enumerate(rows[1:], start=2):
        text = row[col_index["Question"]]
        if text is None:
            continue

        options = {}
        for letter in ["A", "B", "C", "D", "E", "F"]:
            col_name = f"Option_{letter}"
            if col_name in col_index:
                value = row[col_index[col_name]]
                if value is not None:
                    options[letter] = str(value)

        question_type = str(row[col_index["Question_Type"]]).strip().lower()
        correct_raw = str(row[col_index["Correct_Answer"]]).strip()

        if question_type == "multiple":
            correct_answers = [
                a.strip().upper() for a in correct_raw.split(",")
            ]
        else:
            correct_answers = [correct_raw.upper()]

        questions.append(
            Question(
                id=row_num - 1,
                text=str(text),
                options=options,
                correct_answers=correct_answers,
                question_type=question_type,
            )
        )

    return questions


def _parse_alt_format(headers: list, rows: list) -> list[Question]:
    """Parse the alternative XLSX format (e.g. Google Sheets with AI formulas).

    Columns: Question Number, Question Text, Option A, Option B, Option C,
    Option D, Option E (optional), Correct answer (letter + explanation).
    """
    col_index = {name: idx for idx, name in enumerate(headers)}
    questions = []

    for row_num, row in enumerate(rows[1:], start=2):
        text = row[col_index["Question Text"]]
        if text is None:
            continue

        options = {}
        for letter in ["A", "B", "C", "D", "E", "F"]:
            col_name = f"Option {letter}"
            if col_name in col_index:
                value = row[col_index[col_name]]
                if value is not None:
                    options[letter] = str(value)

        correct_raw = row[col_index["Correct answer"]]
        if correct_raw is None:
            correct_raw = ""
        correct_answers, question_type = _parse_alt_correct_answer(str(correct_raw))

        questions.append(
            Question(
                id=row_num - 1,
                text=str(text),
                options=options,
                correct_answers=correct_answers,
                question_type=question_type,
            )
        )

    return questions
